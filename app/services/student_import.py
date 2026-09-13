"""学生名单 Excel 导入：先整批校验，全部通过才落库。

规则（按评审确认）：
1. 表头必须是模板列名；学号、姓名为必填列；
2. 文件内学号重复、学号已存在、姓名为空 —— 任一命中即整批拒绝并返回逐行原因；
3. 校验通过后，建学生账号（默认密码）→ 入班。

分组：模板里**不含**分组序号列（教师流程是建班 → 导名单 → 再建分组 → 再加学生）。
如果历史文件带了「分组序号」列，仍然认：该分组已存在就顺手分进去，还没建就忽略并在
结果里计入 ``ignored_group_hints``，不会因为分组不存在而整批失败。
"""

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlmodel.ext.asyncio.session import AsyncSession

from app.core.exceptions import BusinessRuleError
from app.crud.account import UserRepository
from app.crud.organization import (
    ClassGroupRepository,
    ClassStudentGroupRepository,
    ClassStudentRepository,
)
from app.models.account import SysUser
from app.models.organization import ClassGroup
from app.schemas.organization import StudentImportError, StudentImportResult
from app.services.password import default_password_hash

#: 导入模板的表头（第一行）与对应字段
TEMPLATE_FIELDS: dict[str, str] = {
    "学号": "user_no",
    "姓名": "real_name",
    "专业": "major_name",
    "手机号": "phone",
    "邮箱": "email",
}
TEMPLATE_COLUMNS: tuple[str, ...] = tuple(TEMPLATE_FIELDS)
#: 兼容历史模版的可选列（不在模板里生成，但文件里带了就认）
OPTIONAL_FIELDS: dict[str, str] = {"分组序号": "group_no"}
REQUIRED_COLUMNS: tuple[str, ...] = ("学号", "姓名")
SAMPLE_ROWS: tuple[tuple[object, ...], ...] = (
    ("2026001", "张三", "人工智能技术应用", "13800000000", "zhangsan@example.com"),
    ("2026002", "李四", "人工智能技术应用", "", ""),
)


@dataclass
class StudentRow:
    """Excel 里的一行学生数据。"""

    row_no: int
    user_no: str
    real_name: str
    major_name: str | None = None
    phone: str | None = None
    email: str | None = None
    group_no: int | None = None


def _text(value: object) -> str:
    """单元格转字符串：None / 数字 / 日期统一成去空白的文本。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_students(content: bytes) -> tuple[list[StudentRow], list[StudentImportError]]:
    """解析 xlsx，返回（有效行, 行级格式错误）。表头缺失直接抛业务异常。"""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件抛的异常类型不统一
        raise BusinessRuleError(f"无法解析 Excel 文件：{exc}") from exc

    sheet = workbook.active
    if sheet is None:
        raise BusinessRuleError("Excel 中没有可用的工作表")

    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise BusinessRuleError("Excel 内容为空，请使用导入模板填写")

    header_map = {_text(name): index for index, name in enumerate(header) if _text(name)}
    missing = [column for column in REQUIRED_COLUMNS if column not in header_map]
    if missing:
        raise BusinessRuleError(
            f"缺少必填列：{('、'.join(missing))}；模板列应为：{('、'.join(TEMPLATE_COLUMNS))}"
        )

    parsed: list[StudentRow] = []
    errors: list[StudentImportError] = []
    for offset, raw in enumerate(rows, start=2):
        values: dict[str, str] = {}
        for column, field in {**TEMPLATE_FIELDS, **OPTIONAL_FIELDS}.items():
            index = header_map.get(column)
            values[field] = _text(raw[index]) if index is not None and index < len(raw) else ""
        if not any(values.values()):
            continue  # 整行空白跳过

        group_no: int | None = None
        raw_group = values.get("group_no", "")
        if raw_group:
            try:
                group_no = int(float(raw_group))
            except ValueError:
                errors.append(
                    StudentImportError(row=offset, user_no=values.get("user_no"), reason="分组序号必须是数字")
                )
                continue

        parsed.append(
            StudentRow(
                row_no=offset,
                user_no=values.get("user_no", ""),
                real_name=values.get("real_name", ""),
                major_name=values.get("major_name") or None,
                phone=values.get("phone") or None,
                email=values.get("email") or None,
                group_no=group_no,
            )
        )
    return parsed, errors


def validate_students(
    rows: list[StudentRow],
    *,
    existing_user_nos: set[str],
) -> list[StudentImportError]:
    """整批校验：文件内学号重复、必填缺失、学号已存在。"""
    errors: list[StudentImportError] = []
    seen: dict[str, int] = {}

    for row in rows:
        if not row.user_no:
            errors.append(StudentImportError(row=row.row_no, user_no=None, reason="学号不能为空"))
            continue
        if not row.real_name:
            errors.append(StudentImportError(row=row.row_no, user_no=row.user_no, reason="姓名不能为空"))
        first_seen = seen.get(row.user_no)
        if first_seen is None:
            seen[row.user_no] = row.row_no
        else:
            errors.append(
                StudentImportError(
                    row=row.row_no,
                    user_no=row.user_no,
                    reason=f"学号 {row.user_no} 与第 {first_seen} 行重复",
                )
            )

    for row in rows:
        if row.user_no and row.user_no in existing_user_nos:
            errors.append(
                StudentImportError(
                    row=row.row_no, user_no=row.user_no, reason=f"学号 {row.user_no} 已存在，未导入"
                )
            )
    return errors


async def import_students(
    session: AsyncSession,
    class_id: int,
    rows: list[StudentRow],
    *,
    dry_run: bool = False,
) -> StudentImportResult:
    """把校验通过的行落库（dry_run=True 时只统计不写库）。"""
    users = UserRepository(session)
    groups = ClassGroupRepository(session)
    enrollments = ClassStudentRepository(session)
    memberships = ClassStudentGroupRepository(session)

    enrolled_student_ids = set(await enrollments.student_ids(class_id))
    # 已建的分组（按序号索引）；名单里的分组序号对不上就忽略，不算失败
    class_groups: dict[int, ClassGroup] = {
        group.group_no: group for group in await groups.list_by_class(class_id)
    }

    # 第一遍：只读规划，先把"要建几个账号、要建几条在班记录"算准（dry_run 也走这里）
    created_users = 0
    reused_users = 0
    created_enrollments = 0
    students: dict[str, SysUser | None] = {}
    for row in rows:
        student = await users.get_by(user_no=row.user_no)
        students[row.user_no] = student
        if student is not None:
            reused_users += 1
            if student.id not in enrolled_student_ids:
                created_enrollments += 1
        else:
            created_users += 1
            created_enrollments += 1

    assigned_groups = sum(1 for row in rows if row.group_no in class_groups)
    ignored_group_hints = sum(
        1 for row in rows if row.group_no is not None and row.group_no not in class_groups
    )

    # 第二遍：真正落库（学号重复等已在 validate_students 拦截，这里不再做部分成功）
    if not dry_run:
        for row in rows:
            student = students[row.user_no]
            if student is None:
                student = await users.create(
                    {
                        "user_no": row.user_no,
                        "real_name": row.real_name,
                        "user_type": "STUDENT",
                        "major_name": row.major_name,
                        "phone": row.phone,
                        "email": row.email,
                        "password_hash": default_password_hash(),
                    }
                )
                students[row.user_no] = student

            enrollment = await enrollments.active_enrollment(class_id, student.id)
            if enrollment is None:
                enrollment = await enrollments.create(
                    {"class_id": class_id, "student_id": student.id, "status": "ENROLLED"}
                )

            if row.group_no is None:
                continue
            group = class_groups.get(row.group_no)
            if group is None:
                continue
            await memberships.set_group(enrollment.id, group.id)

    return StudentImportResult(
        dry_run=dry_run,
        class_id=class_id,
        total=len(rows),
        created_users=created_users,
        reused_users=reused_users,
        created_class_students=created_enrollments,
        assigned_groups=assigned_groups,
        ignored_group_hints=ignored_group_hints,
    )


def build_template_xlsx() -> bytes:
    """生成导入模板（表头 + 两行示例）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生名单"
    sheet.append(list(TEMPLATE_COLUMNS))
    for sample in SAMPLE_ROWS:
        sheet.append(list(sample))
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = [
    "OPTIONAL_FIELDS",
    "REQUIRED_COLUMNS",
    "SAMPLE_ROWS",
    "TEMPLATE_COLUMNS",
    "TEMPLATE_FIELDS",
    "StudentRow",
    "build_template_xlsx",
    "import_students",
    "parse_students",
    "validate_students",
]
