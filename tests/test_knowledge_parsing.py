"""解析与切分的纯函数测试（不依赖数据库、MinIO、模型）。"""

import io

import pytest

from app.core.exceptions import BusinessRuleError
from app.services.parsing import looks_binary, parse, sniff_suffix
from app.services.splitting import split_document, strategy_name

# ------------------------------------------------------------------ 解析


def test_parse_markdown_builds_heading_path() -> None:
    content = (
        "# 评分标准\n\n总分 100 分。\n\n## 维度一 需求分析\n\n占 20 分。\n\n"
        "### 细则\n\n需求描述完整得 10 分。\n"
    ).encode()
    document = parse(content, filename="评分标准.md")

    paths = [block.heading_path for block in document.blocks]
    assert paths == ["评分标准", "评分标准 > 维度一 需求分析", "评分标准 > 维度一 需求分析 > 细则"]
    assert document.blocks[0].text == "总分 100 分。"
    assert not document.is_empty


def test_parse_txt_splits_by_blank_line() -> None:
    document = parse("第一段内容。\n\n第二段内容。\n".encode(), filename="说明.txt")

    assert [block.text for block in document.blocks] == ["第一段内容。", "第二段内容。"]
    assert all(block.heading_path == "" for block in document.blocks)


def test_binary_attachment_is_rejected_not_decoded_as_text() -> None:
    """图片/压缩包不能被当成文本"解析成功"。

    gb18030 几乎能解码任意字节序列，只看"解码是否成功"会把 PNG 变成一段乱码正文，
    于是附件被当成有效内容喂给模型、评分标准被当成有效标准入库。必须按文件头挡掉。
    """
    png = b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 4
    assert looks_binary(png) is True
    assert sniff_suffix(png) is None
    with pytest.raises(BusinessRuleError, match="暂不支持解析"):
        parse(png, filename="现场照片.png")

    # 有 NUL 字节的未知二进制同样挡掉
    assert sniff_suffix(b"some header\x00\x01\x02binary") is None
    # 纯文本仍按 markdown 兜底（扩展名丢失时的既有行为）
    assert sniff_suffix("这是一段没有扩展名的中文文本。".encode()) == ".md"


def test_parse_txt_falls_back_to_gbk() -> None:
    document = parse("评分标准：需求分析 20 分。".encode("gb18030"), filename="准则.txt")

    assert "需求分析" in document.blocks[0].text


def test_parse_csv_expands_header_pairs() -> None:
    document = parse("维度,分值\n需求分析,20\n方案设计,15\n".encode(), filename="评分表.csv")

    assert len(document.blocks) == 1
    assert document.blocks[0].text == "维度: 需求分析；分值: 20\n维度: 方案设计；分值: 15"


def test_parse_docx_reads_headings_and_tables() -> None:
    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_heading("评分标准", level=1)
    docx.add_paragraph("总分 100 分。")
    docx.add_heading("维度一", level=2)
    docx.add_paragraph("需求分析占 20 分。")
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "维度"
    table.cell(0, 1).text = "分值"
    table.cell(1, 0).text = "需求分析"
    table.cell(1, 1).text = "20"
    buffer = io.BytesIO()
    docx.save(buffer)

    document = parse(buffer.getvalue(), filename="评分标准.docx")

    assert [block.heading_path for block in document.blocks] == [
        "评分标准",
        "评分标准 > 维度一",
        "评分标准 > 维度一",
    ]
    assert document.blocks[1].text == "需求分析占 20 分。"
    assert document.blocks[2].text == "维度 | 分值\n需求分析 | 20"


def test_parse_xlsx_uses_sheet_name_as_heading() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "评分表"
    sheet.append(["维度", "分值"])
    sheet.append(["需求分析", 20])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    document = parse(buffer.getvalue(), filename="评分表.xlsx")

    assert document.blocks[0].heading_path == "评分表"
    assert "维度: 需求分析；分值: 20" in document.blocks[0].text


def test_parse_pdf_without_text_is_empty() -> None:
    """扫描件（无文本层）解析出来是空的，由入库环节判定失败并给原因。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    document = parse(buffer.getvalue(), filename="扫描件.pdf")

    assert document.blocks == []
    assert document.is_empty


def test_parse_rejects_unsupported_suffix() -> None:
    # 普通 zip（PK 头但既不是 docx 也不是 xlsx）既没扩展名支持、也嗅探不出格式 → 明确失败
    with pytest.raises(BusinessRuleError) as exc:
        parse(b"PK\x03\x04\x14\x00\x00\x00\x08\x00" + "归档内容".encode(), filename="归档.zip")

    assert ".zip" in str(exc.value)


def test_parse_sniffs_suffix_when_filename_lost_extension() -> None:
    """客户端把中文文件名编码坏（丢掉扩展名）时，按文件头猜格式。"""
    document = parse("# 评分标准\n\n需求分析占 20 分。\n".encode(), filename="=?utf-8?B?6K+E5YiG?=")

    assert document.blocks[0].heading_path == "评分标准"


def test_parse_sniffs_docx_by_magic_bytes() -> None:
    from docx import Document as DocxDocument

    docx = DocxDocument()
    docx.add_heading("评分标准", level=1)
    docx.add_paragraph("需求分析占 20 分。")
    buffer = io.BytesIO()
    docx.save(buffer)

    document = parse(buffer.getvalue(), filename="无名文件")

    assert document.blocks[0].heading_path == "评分标准"


def test_clean_upload_name_decodes_rfc2047_and_strips_path() -> None:
    from app.services.storage import clean_upload_name

    assert clean_upload_name("=?utf-8?B?6K+E5YiG5qCH5YeG56S65L6LLm1k?=").endswith(".md")
    assert clean_upload_name("..\\..\\etc\\passwd") == "passwd"
    assert clean_upload_name(None) == "unnamed"


# ------------------------------------------------------------------ 切分


def test_split_keeps_heading_prefix_and_hash() -> None:
    document = parse("# 评分标准\n\n需求分析占 20 分。\n".encode(), filename="准则.md")
    chunks = split_document(document, chunk_size_chars=100, overlap_chars=20, max_chunk_chars=1000)

    assert len(chunks) == 1
    assert chunks[0].content == "评分标准\n需求分析占 20 分。"
    assert chunks[0].heading_path == "评分标准"
    assert chunks[0].char_count == len(chunks[0].content)
    assert len(chunks[0].content_hash) == 64


def test_split_respects_target_size_and_overlap() -> None:
    sentences = "".join(f"第{i}条要求：说明清楚扣分标准，不能含糊。\n" for i in range(1, 21))
    document = parse(sentences.encode(), filename="准则.txt")
    chunks = split_document(document, chunk_size_chars=80, overlap_chars=20, max_chunk_chars=1000)

    assert len(chunks) > 3
    assert all(chunk.char_count <= 1000 for chunk in chunks)
    # 相邻块重叠：前一块的尾部应出现在后一块开头
    for previous, current in zip(chunks, chunks[1:], strict=False):
        tail = previous.content[-20:].strip()
        assert tail and tail[:10] in current.content[:40]


def test_split_hard_caps_oversized_single_sentence() -> None:
    document = parse(("啊" * 500).encode(), filename="长句.txt")
    chunks = split_document(document, chunk_size_chars=100, overlap_chars=10, max_chunk_chars=120)

    assert len(chunks) >= 4
    assert all(chunk.char_count <= 120 for chunk in chunks)
    assert "".join(chunk.content for chunk in chunks).count("啊") >= 500


def test_split_separates_by_page_for_pdf_like_blocks() -> None:
    from app.services.parsing import ParsedBlock, ParsedDocument

    document = ParsedDocument(
        blocks=[
            ParsedBlock(text="第一页内容。", page_no=1),
            ParsedBlock(text="第二页内容。", page_no=2),
        ]
    )
    chunks = split_document(document, chunk_size_chars=100, overlap_chars=10, max_chunk_chars=1000)

    assert [chunk.page_no for chunk in chunks] == [1, 2]


def test_whole_document_split_keeps_headings_and_makes_one_chunk() -> None:
    """整份切分：一份标准就是一块，且分节标题要回到正文里。

    解析阶段把 ``##`` 收进了 heading_path，正文里没有；整份切分不还原标题的话，
    模型拿到的标准就没有"一、需求分析（10 分）"这类分节名，没法把条款和关卡对应起来。
    """
    content = (
        "# 评分标准\n\n总分 100 分。\n\n"
        "## 一、需求分析（10 分）\n\n必须覆盖六项，每缺一项扣 1.5 分。\n\n"
        "## 二、方案设计（15 分）\n\n只写结论每项扣 1 分。\n"
    ).encode()
    document = parse(content, filename="评分标准.md")

    chunks = split_document(document, whole=True)
    assert len(chunks) == 1
    text = chunks[0].content
    assert text.startswith("## 评分标准")  # 一级标题也还原回来了
    assert "一、需求分析（10 分）" in text
    assert "二、方案设计（15 分）" in text
    assert "必须覆盖六项" in text and "只写结论每项扣 1 分" in text
    assert chunks[0].heading_path == ""
    assert chunks[0].char_count == len(text)

    # 同一份内容按结构化切分仍是多片——两条路径互不影响
    assert len(split_document(document)) == 3


def test_build_plan_routes_strategy_by_doc_type() -> None:
    """评分标准整份一块；知识库资料仍然结构化细切。"""
    from app.models.enums import KnowledgeDocType
    from app.services.knowledge_ingest import build_plan

    content = (
        "# 评分标准\n\n## 一、需求分析（10 分）\n\n必须覆盖六项。\n\n"
        "## 二、方案设计（15 分）\n\n必须给出推导。\n"
    ).encode()

    criteria = build_plan(content, "标准.md", doc_type=KnowledgeDocType.EVAL_CRITERIA)
    assert criteria.strategy == "whole_document"
    assert len(criteria.chunks) == 1

    knowledge = build_plan(content, "资料.md", doc_type=KnowledgeDocType.KNOWLEDGE)
    assert knowledge.strategy == "v1_structural_800"
    assert len(knowledge.chunks) == 2  # 两个小节各自成片


def test_content_hash_ignores_whitespace_difference() -> None:
    first = split_document(parse("需求分析  占   20 分。".encode(), filename="a.txt"), chunk_size_chars=100)
    second = split_document(parse("需求分析 占 20 分。".encode(), filename="b.txt"), chunk_size_chars=100)

    assert first[0].content_hash == second[0].content_hash
    assert strategy_name(800) == "v1_structural_800"
